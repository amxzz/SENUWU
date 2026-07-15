import os
import json
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def add_toc(paragraph):
    run = paragraph.add_run()
    fldChar = OxmlElement('w:fldChar')
    fldChar.set(qn('w:fldCharType'), 'begin')
    
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = 'TOC \\o "1-2" \\h \\z \\u'
    
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    
    r_element = run._r
    r_element.append(fldChar)
    r_element.append(instrText)
    r_element.append(fldChar2)
    r_element.append(fldChar3)

def configure_styles(doc):
    # Configure Normal
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    paragraph_format = style.paragraph_format
    paragraph_format.line_spacing = 1.5
    paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph_format.first_line_indent = Inches(0.5)

    # Configure Heading 1
    style_h1 = doc.styles['Heading 1']
    font_h1 = style_h1.font
    font_h1.name = 'Times New Roman'
    font_h1.size = Pt(14)
    font_h1.bold = True
    style_h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_h1.paragraph_format.line_spacing = 1.5

    # Configure Heading 2
    style_h2 = doc.styles['Heading 2']
    font_h2 = style_h2.font
    font_h2.name = 'Times New Roman'
    font_h2.size = Pt(12)
    font_h2.bold = True
    style_h2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    style_h2.paragraph_format.line_spacing = 1.5

def generate_report():
    template_path = r"D:\Kuliyeah\sem 6\Project HKI SENUWU\TEMPLATE LAPORAN PBL.docx"
    output_path = r"D:\Kuliyeah\sem 6\Project HKI SENUWU\LAPORAN PBLL .docx"
    assets_dir = r"D:\Kuliyeah\sem 6\Project HKI SENUWU\assets"
    sumber_dir = r"D:\Kuliyeah\sem 6\Project HKI SENUWU\sumber"
    
    # Load metadata
    metadata_path = os.path.join(sumber_dir, "references.json")
    references = []
    if os.path.exists(metadata_path):
        with open(metadata_path, 'r', encoding='utf-8') as f:
            references = json.load(f)
            
    doc = Document(template_path)
    configure_styles(doc)

    # Determine citations for each chapter (distributing the 15+ refs)
    def get_cites(start, count):
        cites = []
        for i in range(start, min(start+count, len(references))):
            ref = references[i]
            author = ref['authors'][0].split()[-1] if ref['authors'] else "Anon"
            cites.append(f"({author} et al., {ref['year']})")
        return ", ".join(cites)

    # Build Content
    doc.add_page_break()
    doc.add_heading('DAFTAR ISI', level=1)
    p_toc = doc.add_paragraph()
    add_toc(p_toc)
    doc.add_page_break()

    # BAB 1
    doc.add_heading('BAB I PENDAHULUAN', level=1)
    doc.add_heading('1.1 Latar Belakang', level=2)
    cites_bab1 = get_cites(0, 4)
    doc.add_paragraph(f"Ketimpangan spasial merupakan salah satu tantangan terbesar dalam perencanaan pembangunan di Indonesia. Untuk mengatasi permasalahan ini, analisis berbasis data yang akurat sangat dibutuhkan {cites_bab1}. Sistem Analitik SENUWU hadir sebagai solusi inovatif dalam bentuk Decision Support System (DSS).")
    doc.add_paragraph("Target pengguna dan dampak utama dari sistem ini dirancang secara khusus untuk Pemerintah (seperti Bappenas dan Pemerintah Daerah) sebagai instrumen Decision Support System (DSS) untuk mendukung kebijakan pembangunan yang merata. Selain itu, sistem ini juga berfungsi sebagai platform transparansi dan pengetahuan publik mengenai kondisi kesejahteraan di berbagai wilayah.")
    
    doc.add_heading('1.2 Tujuan', level=2)
    doc.add_paragraph("Tujuan utama dari sistem ini adalah menyediakan alat klasterisasi kesejahteraan spasial yang akurat dan *robust* terhadap data anomali.")
    
    # BAB 2
    doc.add_heading('BAB II TINJAUAN PUSTAKA', level=1)
    doc.add_heading('2.1 Pendekatan Algoritma Hibrida', level=2)
    cites_bab2 = get_cites(4, 4)
    doc.add_paragraph(f"Klasterisasi spasial seringkali terkendala oleh adanya noise dan outlier. Penggunaan algoritma konvensional seperti K-Means dapat menghasilkan centroid yang bias. Oleh karena itu, pendekatan hibrida diperlukan {cites_bab2}. Nilai kebaruan (Novelty) yang berpotensi menjadi Hak Kekayaan Intelektual (HKI) dari sistem SENUWU ini terletak pada pendekatan Hibrida yang secara berurutan menggabungkan algoritma Isolation Forest, HDBSCAN, dan K-Means.")
    
    doc.add_heading('2.2 Validasi Klaster', level=2)
    cites_bab2_2 = get_cites(8, 3)
    doc.add_paragraph(f"Untuk menguji keandalan hasil klasterisasi, digunakan metrik evaluasi yang disesuaikan dengan konteks kesejahteraan. {cites_bab2_2}. Sistem ini memperkenalkan indeks validasi spesifik yaitu Welfare Clustering Validity Index (WCVI) yang dioptimasi untuk mengukur rasio ketimpangan spasial dengan lebih akurat.")
    
    # BAB 3
    doc.add_heading('BAB III SPESIFIKASI DAN ARSITEKTUR SISTEM', level=1)
    doc.add_heading('3.1 Arsitektur Sistem', level=2)
    cites_bab3 = get_cites(11, 4)
    doc.add_paragraph(f"Sistem dibangun berbasis web menggunakan framework Streamlit untuk memastikan interaktivitas yang tinggi dan performa visualisasi peta yang optimal {cites_bab3}.")
    
    # Insert Extracted Image (Architecture)
    img_arch = os.path.join(assets_dir, "img_p8_1.jpeg")
    if os.path.exists(img_arch):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_img = p_img.add_run()
        run_img.add_picture(img_arch, width=Inches(5.0))
        p_cap = doc.add_paragraph("Gambar 3.1 Arsitektur Data Pipeline SENUWU")
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph("[... Kebutuhan Visual: Diagram arsitektur data preprocessing ketimpangan spasial yang merinci tahapan cleaning dan imputasi missing values secara lebih granular ...]")
    
    doc.add_heading('3.2 Antarmuka Pengguna (UI)', level=2)
    img_ui = os.path.join(assets_dir, "img_p9_1.png")
    if os.path.exists(img_ui):
        p_img2 = doc.add_paragraph()
        p_img2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_img2 = p_img2.add_run()
        run_img2.add_picture(img_ui, width=Inches(5.5))
        p_cap2 = doc.add_paragraph("Gambar 3.2 Antarmuka Peta Choropleth SENUWU")
        p_cap2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
    img_ui2 = os.path.join(assets_dir, "img_p10_1.jpeg")
    if os.path.exists(img_ui2):
        p_img3 = doc.add_paragraph()
        p_img3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_img3 = p_img3.add_run()
        run_img3.add_picture(img_ui2, width=Inches(5.5))
        p_cap3 = doc.add_paragraph("Gambar 3.3 Antarmuka Detail Hasil Klasterisasi")
        p_cap3.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # BAB 4
    doc.add_heading('BAB IV KESIMPULAN', level=1)
    doc.add_heading('4.1 Kesimpulan', level=2)
    doc.add_paragraph("Sistem Analitik SENUWU berhasil mengintegrasikan model Hibrida (Isolation Forest + HDBSCAN + K-Means) dengan WCVI untuk menghasilkan peta ketimpangan yang andal. Sistem ini siap digunakan oleh Bappenas dan Pemerintah Daerah sebagai instrumen Decision Support System yang komprehensif.")

    # DAFTAR PUSTAKA
    doc.add_page_break()
    doc.add_heading('DAFTAR PUSTAKA', level=1)
    
    style_bib = doc.styles.add_style('Bibliography', 1)
    style_bib.font.name = 'Times New Roman'
    style_bib.font.size = Pt(12)
    style_bib.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    style_bib.paragraph_format.left_indent = Inches(0.5)
    style_bib.paragraph_format.first_line_indent = Inches(-0.5) # Hanging indent
    style_bib.paragraph_format.line_spacing = 1.0
    
    for ref in references:
        authors_str = ", ".join(ref['authors']) if ref['authors'] else "Anon"
        bib_text = f"{authors_str}. ({ref['year']}). {ref['title']}. Diakses dari arXiv."
        doc.add_paragraph(bib_text, style='Bibliography')

    doc.save(output_path)
    print(f"Laporan berhasil dibuat dan disimpan di: {output_path}")

if __name__ == "__main__":
    generate_report()
